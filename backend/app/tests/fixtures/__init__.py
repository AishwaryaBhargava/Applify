"""Test fixtures: a fictional resume plus in-memory PDF and DOCX builders.

Nothing here touches the network or the disk. The PDF is assembled by hand
rather than with a rendering library so the test suite needs no extra
dependency, and the DOCX is built with python-docx, which is already a runtime
dependency of the parser it tests.

The resume is entirely fictional. No real person's details appear in this repo.
"""

import io

from docx import Document

# --------------------------------------------------------------------------
# A realistic, fictional resume
# --------------------------------------------------------------------------

# Written to exercise every section the extractor is asked for: a summary, two
# roles (one current), two degrees, a grouped skills line that has to be split
# into tokens, certifications, projects with links, and standalone achievements.
SAMPLE_RESUME_TEXT = """\
JORDAN A. RIVERA
Bengaluru, India | jordan.rivera@example.com | +91 90000 00000
linkedin.com/in/jordanrivera-example | github.com/jordanrivera-example

PROFESSIONAL SUMMARY
Backend engineer with six years building payment and identity services at high
transaction volume. Comfortable owning a service end to end, from schema design
through on-call. Most effective on teams that ship incrementally and measure
what they ship.

WORK EXPERIENCE

Senior Backend Engineer
Kestrel Payments, Bengaluru, India
March 2022 - Present
- Led the migration of the settlement ledger from a monolith to three Go
  services, cutting end-of-day reconciliation from 40 minutes to 6.
- Designed an idempotency layer for the payouts API that eliminated duplicate
  disbursements, previously averaging 12 per month.
- Introduced contract tests across 9 internal services, reducing integration
  regressions caught in staging by 60%.
- Mentored 4 engineers; 2 were promoted within 18 months.

Backend Engineer
Northwind Identity, Pune, India
July 2019 - February 2022
- Built the SSO provisioning service in Python and FastAPI, onboarding 140
  enterprise tenants.
- Cut p99 token introspection latency from 380ms to 74ms by adding a Redis
  read-through cache and reworking the session index.
- Owned the PostgreSQL schema for the directory sync pipeline, processing 2.5M
  records nightly.

EDUCATION

Master of Technology, Computer Science
Indian Institute of Technology, Hyderabad
2017 - 2019
CGPA 8.7/10. Thesis on consistency models for distributed key-value stores.

Bachelor of Engineering, Information Technology
Pune Institute of Technology
2013 - 2017
First Class with Distinction.

SKILLS
Languages: Python, Go, SQL, TypeScript
Frameworks: FastAPI, Django, gRPC
Data: PostgreSQL, Redis, Kafka, ClickHouse
Infrastructure: Docker, Kubernetes, Terraform, AWS
Practices: Distributed systems design, Observability, Incident response

CERTIFICATIONS
AWS Certified Solutions Architect - Associate, Amazon Web Services, 2021
Certified Kubernetes Administrator, Cloud Native Computing Foundation, 2023

PROJECTS

Ledgerlite
An append-only double-entry ledger library with a SQLite and PostgreSQL
backend. Used by two small fintech teams in production.
Technologies: Go, PostgreSQL, SQLite
github.com/jordanrivera-example/ledgerlite

Tracewire
A minimal OpenTelemetry collector that samples traces by error budget burn
rather than by fixed rate.
Technologies: Python, OpenTelemetry, Prometheus

ACHIEVEMENTS
- Speaker, IndiaFOSS 2023: "Idempotency is a Product Decision".
- Winner, Kestrel Payments internal hackathon 2022, for a fraud replay tool.
- Published "Reconciling at Scale" in the Kestrel engineering blog, 2023.

References available on request.
"""

# --------------------------------------------------------------------------
# In-memory document builders
# --------------------------------------------------------------------------


def _escape_pdf_text(value: str) -> str:
    """Escape the three characters that are special inside a PDF string."""
    return value.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def make_pdf(text: str, font_size: int = 11) -> bytes:
    """Build a single-page PDF whose text pdfplumber can read back.

    Hand-rolled rather than rendered with reportlab: the parser under test only
    needs a structurally valid PDF with a Type1 base font, and building one here
    keeps a rendering library out of requirements.txt for the sake of a fixture.

    Args:
        text: The document body. Each line becomes one line on the page.
        font_size: Point size, which also sets the leading.

    Returns:
        The PDF file's bytes.
    """
    lines = text.splitlines() or [""]

    content: list[str] = [
        "BT",
        "/F1 {} Tf".format(font_size),
        "{} TL".format(font_size + 3),
        "50 750 Td",
    ]
    for index, line in enumerate(lines):
        if index:
            content.append("T*")  # move down one leading
        content.append("({}) Tj".format(_escape_pdf_text(line)))
    content.append("ET")
    # latin-1 is what a PDF Type1 base font expects for its byte strings.
    stream = "\n".join(content).encode("latin-1", errors="replace")

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(str(number).encode("ascii") + b" 0 obj\n" + body + b"\nendobj\n")

    # The cross-reference table has to record each object's byte offset, so it
    # can only be written once the body is laid out.
    xref_offset = out.tell()
    size = len(objects) + 1
    out.write(b"xref\n0 " + str(size).encode("ascii") + b"\n")
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write("{:010d} 00000 n \n".format(offset).encode("ascii"))
    out.write(
        b"trailer\n<< /Size "
        + str(size).encode("ascii")
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_offset).encode("ascii")
        + b"\n%%EOF\n"
    )
    return out.getvalue()


def make_docx(text: str, table_rows: list[list[str]] | None = None) -> bytes:
    """Build a DOCX with one paragraph per line, plus an optional table.

    Args:
        text: The document body. Each line becomes one paragraph.
        table_rows: Rows of cell text appended as a table, so the parser's
            table handling can be exercised.

    Returns:
        The DOCX file's bytes.
    """
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
