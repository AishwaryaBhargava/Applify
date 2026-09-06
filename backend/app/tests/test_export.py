"""Tests for Markdown-to-DOCX rendering and the export route.

The renderer is checked by *reading the file back*: every assertion opens the
returned bytes with python-docx and inspects the paragraphs, styles, runs, and
XML that Word will see. Asserting on the input Markdown would prove nothing --
the whole risk in this module is that a file is produced which Word refuses to
open or renders as one undifferentiated block.

The route is checked against the ``FakeSession``: filenames, both formats, and
the 404 that keeps one user's generated documents out of another's downloads.
"""

import uuid
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient

from app.api.routes.outputs import OUTPUT_NOT_FOUND
from app.services import export_service
from app.services.export_service import (
    export_filename,
    markdown_to_docx,
    slugify,
)
from app.tests.conftest import OTHER_USER_ID, FakeSession
from app.tests.test_chat import make_chat, make_tracker
from app.tests.test_outputs import make_output

CHAT_ID = "00000000-0000-0000-0000-000000000001"

RESUME_MARKDOWN = """# Ada Lovelace

ada@example.com | London

## Summary

A **backend** engineer with *six* years in payments.

---

## Experience

### Senior Engineer, Kestrel Payments

- Cut reconciliation from 40 minutes to 6.
* Owned the CI/CD pipeline.

## Links

1. Portfolio
2. [My site](https://ada.example)
"""

# The MIME type Word registers itself for. Spelled out here so a typo in the
# service is a failing test rather than a file the browser saves as .zip.
DOCX_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def render(markdown: str = RESUME_MARKDOWN, title: str = "Resume"):
    """Render Markdown and hand back the reopened Word document."""
    return Document(BytesIO(markdown_to_docx(markdown, title)))


def paragraph_named(document, text: str):
    """Return the first paragraph whose text matches exactly."""
    for paragraph in document.paragraphs:
        if paragraph.text == text:
            return paragraph
    raise AssertionError("no paragraph reads {!r}".format(text))


def has_bottom_rule(paragraph) -> bool:
    """Whether this paragraph carries an explicit bottom border."""
    properties = paragraph._p.pPr
    if properties is None:
        return False
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    borders = properties.find(namespace + "pBdr")
    return borders is not None and borders.find(namespace + "bottom") is not None


def seeded_output(db: FakeSession, **kwargs):
    """Seed a chat, its tracker entry, and one generated output."""
    chat_kwargs = {"company": kwargs.pop("company", "Kestrel Payments")}
    if "user_id" in kwargs:
        chat_kwargs["user_id"] = kwargs.pop("user_id")
    chat = db.seed(make_chat(**chat_kwargs))
    db.seed(make_tracker(chat.id, user_id=str(chat.user_id)))
    output = db.seed(make_output(chat.id, **kwargs))
    return chat, output


# ==========================================================================
# Auth guard
# ==========================================================================


def test_export_requires_auth(client: TestClient) -> None:
    """GET /chats/{id}/outputs/{id}/export is protected by get_current_user."""
    response = client.get(
        "/chats/{}/outputs/{}/export".format(CHAT_ID, uuid.uuid4())
    )
    assert response.status_code == 401


# ==========================================================================
# Document setup
# ==========================================================================


def test_the_page_is_set_up_the_way_a_resume_should_be() -> None:
    """Margins, typeface, size, and single spacing, all as configured."""
    document = render()
    normal = document.styles["Normal"]

    assert document.sections[0].left_margin.inches == 0.75
    assert document.sections[0].top_margin.inches == 0.75
    assert normal.font.name == "Calibri"
    assert normal.font.size == Pt(11)
    assert normal.paragraph_format.line_spacing == 1.0


def test_the_title_is_stored_as_a_document_property() -> None:
    """The file names itself in a document library, without printing a heading."""
    document = render(title="Kestrel Payments - resume")

    assert document.core_properties.title == "Kestrel Payments - resume"
    assert "Kestrel Payments - resume" not in [
        p.text for p in document.paragraphs
    ]


# ==========================================================================
# Headings
# ==========================================================================


def test_h1_uses_the_title_style_at_sixteen_point() -> None:
    """The candidate's name is the document's title, in the accent colour."""
    heading = paragraph_named(render(), "Ada Lovelace")
    run = heading.runs[0]

    assert heading.style.name == "Title"
    assert run.bold is True
    assert run.font.size == Pt(16)
    assert run.font.color.rgb == export_service.TITLE_COLOR


def test_h2_is_twelve_point_bold_with_a_rule_under_it() -> None:
    """Section headers are drawn, not borrowed from Word's blue Heading 2."""
    heading = paragraph_named(render(), "Summary")

    assert heading.runs[0].bold is True
    assert heading.runs[0].font.size == Pt(12)
    assert has_bottom_rule(heading) is True


def test_h3_is_eleven_point_bold_with_no_rule() -> None:
    """Role headings sit inside a section, so they get weight but no divider."""
    heading = paragraph_named(render(), "Senior Engineer, Kestrel Payments")

    assert heading.runs[0].bold is True
    assert heading.runs[0].font.size == Pt(11)
    assert has_bottom_rule(heading) is False


def test_deep_headings_collapse_onto_the_third_level() -> None:
    """An H5 in a resume is a model slip, not a fourth kind of heading."""
    heading = paragraph_named(render("##### Deep\n"), "Deep")

    assert heading.runs[0].font.size == Pt(11)


def test_the_border_is_placed_where_the_schema_requires() -> None:
    """pBdr after pStyle and before spacing -- Word calls any other order corrupt."""
    heading = paragraph_named(render(), "Summary")
    tags = [child.tag.split("}")[1] for child in heading._p.pPr]

    assert tags.index("pBdr") < tags.index("spacing")


# ==========================================================================
# Body
# ==========================================================================


def test_bullets_use_the_list_bullet_style_whichever_marker_was_used() -> None:
    """"-" and "*" are the same list to a reader, so they are here too."""
    document = render()

    assert (
        paragraph_named(
            document, "Cut reconciliation from 40 minutes to 6."
        ).style.name
        == "List Bullet"
    )
    assert paragraph_named(document, "Owned the CI/CD pipeline.").style.name == (
        "List Bullet"
    )


def test_numbered_items_use_the_list_number_style() -> None:
    """A numbered list keeps its numbering to Word, not as literal "1." text."""
    assert paragraph_named(render(), "Portfolio").style.name == "List Number"


def test_bold_and_italic_become_runs_not_asterisks() -> None:
    """The markers are consumed; the emphasis survives as formatting."""
    paragraph = paragraph_named(
        render(), "A backend engineer with six years in payments."
    )
    runs = {run.text: run for run in paragraph.runs}

    assert runs["backend"].bold is True
    assert runs["six"].italic is True
    assert "*" not in paragraph.text


def test_unmatched_emphasis_survives_as_literal_text() -> None:
    """A malformed document still exports; it just keeps the stray asterisks."""
    paragraph = paragraph_named(render("An **unclosed line\n"), "An **unclosed line")

    assert paragraph.text == "An **unclosed line"


def test_links_are_flattened_to_text_and_url() -> None:
    """A hyperlink field is invisible on paper and to a plain-text ATS parse."""
    assert paragraph_named(render(), "My site (https://ada.example)")


def test_a_bare_link_is_not_repeated() -> None:
    """"url (url)" helps nobody."""
    document = render("[https://ada.example](https://ada.example)\n")

    assert document.paragraphs[0].text == "https://ada.example"


def test_a_horizontal_rule_becomes_blank_space() -> None:
    """A printed rule across a resume reads as a formatting accident."""
    document = render()
    texts = [p.text for p in document.paragraphs]

    assert "---" not in texts
    assert "" in texts


def test_blank_lines_do_not_become_empty_paragraphs() -> None:
    """Markdown's blank lines are separators; the spacing is set by the style."""
    document = render("One\n\n\n\nTwo\n")

    assert [p.text for p in document.paragraphs] == ["One", "Two"]


def test_an_empty_document_still_renders() -> None:
    """Exporting an output that somehow has no content is a file, not a 500."""
    document = render("", title="Empty")

    assert document.paragraphs == [] or all(
        not p.text for p in document.paragraphs
    )


# ==========================================================================
# Filenames
# ==========================================================================


@pytest.mark.parametrize(
    "text, expected",
    [
        ("Kestrel Payments", "kestrel-payments"),
        ("Zürich Insurance", "zurich-insurance"),
        ("  A&B  Co.  ", "a-b-co"),
        ("!!!", "applify"),
        ("", "applify"),
    ],
)
def test_slugify(text: str, expected: str) -> None:
    """A readable, safe fragment -- or the fallback, never an empty name."""
    assert slugify(text) == expected


def test_export_filename_names_the_company_and_the_kind() -> None:
    """The two things that tell twenty downloads apart six weeks later."""
    assert (
        export_filename("Kestrel Payments", "cover_letter", "docx")
        == "kestrel-payments-cover-letter.docx"
    )
    assert export_filename("Kestrel", "resume", "md") == "kestrel-resume.md"
    assert export_filename(None, "answer", "docx") == "applify-answer.docx"


# ==========================================================================
# GET /chats/{id}/outputs/{output_id}/export
# ==========================================================================


def test_export_returns_a_docx_attachment(
    auth_client: TestClient, db: FakeSession
) -> None:
    """The default format is the one people upload to application forms."""
    chat, output = seeded_output(db, content=RESUME_MARKDOWN)

    response = auth_client.get(
        "/chats/{}/outputs/{}/export".format(chat.id, output.id)
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="kestrel-payments-resume.docx"'
    )
    assert paragraph_named(Document(BytesIO(response.content)), "Ada Lovelace")


def test_export_names_the_file_after_the_document_kind(
    auth_client: TestClient, db: FakeSession
) -> None:
    """cover_letter becomes cover-letter, not cover_letter."""
    chat, output = seeded_output(db, output_type="cover_letter", content="Dear team")

    response = auth_client.get(
        "/chats/{}/outputs/{}/export".format(chat.id, output.id)
    )

    assert 'filename="kestrel-payments-cover-letter.docx"' in (
        response.headers["content-disposition"]
    )


def test_export_falls_back_to_the_job_title_without_a_company(
    auth_client: TestClient, db: FakeSession
) -> None:
    """A chat with no company still produces a filename someone can read."""
    chat, output = seeded_output(db, company=None)

    response = auth_client.get(
        "/chats/{}/outputs/{}/export".format(chat.id, output.id)
    )

    assert 'filename="backend-engineer-resume.docx"' in (
        response.headers["content-disposition"]
    )


def test_export_can_return_the_raw_markdown(
    auth_client: TestClient, db: FakeSession
) -> None:
    """format=md hands back exactly what was generated, byte for byte."""
    chat, output = seeded_output(db, content=RESUME_MARKDOWN)

    response = auth_client.get(
        "/chats/{}/outputs/{}/export?format=md".format(chat.id, output.id)
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/markdown")
    assert 'filename="kestrel-payments-resume.md"' in (
        response.headers["content-disposition"]
    )
    assert response.text == RESUME_MARKDOWN


def test_export_rejects_an_unknown_format(
    auth_client: TestClient, db: FakeSession
) -> None:
    """PDF is not offered, so asking for one is a 422 rather than a docx."""
    chat, output = seeded_output(db)

    response = auth_client.get(
        "/chats/{}/outputs/{}/export?format=pdf".format(chat.id, output.id)
    )

    assert response.status_code == 422


def test_export_404s_on_another_users_chat(
    auth_client: TestClient, db: FakeSession
) -> None:
    """Someone else's generated resume is not downloadable, and does not exist."""
    chat, output = seeded_output(db, user_id=OTHER_USER_ID)

    response = auth_client.get(
        "/chats/{}/outputs/{}/export".format(chat.id, output.id)
    )

    assert response.status_code == 404


def test_export_404s_on_an_output_from_a_different_chat(
    auth_client: TestClient, db: FakeSession
) -> None:
    """An output id borrowed from another chat does not resolve inside this one."""
    mine, _ = seeded_output(db)
    _, theirs = seeded_output(db, user_id=OTHER_USER_ID)

    response = auth_client.get(
        "/chats/{}/outputs/{}/export".format(mine.id, theirs.id)
    )

    assert response.status_code == 404
    assert response.json()["detail"] == OUTPUT_NOT_FOUND


def test_export_404s_on_an_unknown_output(
    auth_client: TestClient, db: FakeSession
) -> None:
    """An output id that was never issued is a 404, not a 500."""
    chat, _ = seeded_output(db)

    response = auth_client.get(
        "/chats/{}/outputs/{}/export".format(chat.id, uuid.uuid4())
    )

    assert response.status_code == 404
